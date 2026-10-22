"""Editable PowerPoint, PDF handout, and previews from the same drawing primitives."""
import textwrap
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

BG, PANEL, WHITE, MUTED, CYAN = "10243B", "19344F", "F5F8FC", "BDD0E2", "55DECA"


class Slide:
    """Mirror PowerPoint elements to a PDF page; positions are in inches."""
    def __init__(self, prs, number, title, section, pdf):
        self.slide = prs.slides.add_slide(prs.slide_layouts[6])
        self.slide.background.fill.solid()
        self.slide.background.fill.fore_color.rgb = RGBColor.from_string(BG)
        self.fig = plt.figure(figsize=(13.333, 7.5), facecolor="#" + BG)
        self.ax = self.fig.add_axes([0, 0, 1, 1])
        self.ax.set(xlim=(0, 13.333), ylim=(7.5, 0)); self.ax.axis("off")
        self.pdf = pdf
        self.rect(0, 0, 13.333, .06, CYAN)
        self.text(section.upper(), .6, .32, 12, .3, 10, CYAN, True)
        self.text(title, .6, .84, 12.1, .85, 28, WHITE, True)
        self.text("INSURANCE BEHAVIOR LAB   /   DATA SINTETIS   /   WORKSHOP", .6, 7.05, 11, .22, 9, MUTED)
        self.text(f"{number:02d}", 12.1, 7.05, .6, .22, 10, CYAN, True)

    def rect(self, x, y, w, h, color=PANEL):
        shape = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(color)
        shape.line.fill.background()
        self.ax.add_patch(plt.Rectangle((x, y), w, h, color="#" + color, linewidth=0))

    def text(self, value, x, y, w, h, size=17, color=WHITE, bold=False):
        textbox = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = textbox.text_frame
        frame.word_wrap = False
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        # Explicit wrapping keeps PowerPoint and companion PDF aligned.
        width = max(10, int(w * 72 / (size * .54)))
        lines = []
        for paragraph in str(value).split("\n"):
            lines.extend(textwrap.wrap(paragraph, width=width, break_long_words=False) or [""])
        for i, line in enumerate(lines):
            p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            p.text = line; p.font.name = "DejaVu Sans"; p.font.size = Pt(size)
            p.font.bold = bold; p.font.color.rgb = RGBColor.from_string(color)
            p.space_before = Pt(0); p.space_after = Pt(0); p.line_spacing = 1.18
        line_height = size * 1.18 / 72
        if len(lines) * line_height > h + .07:
            raise ValueError(f"Text overflow: {value[:70]} ({len(lines)} lines in {h} inches)")
        self.ax.text(x, y, "\n".join(lines), fontsize=size, color="#" + color,
                     weight="bold" if bold else "normal", va="top", linespacing=1.18,
                     fontfamily="DejaVu Sans")

    def image(self, path, x=.65, y=1.86, w=12.0, h=4.65):
        img = plt.imread(path)
        aspect = img.shape[1] / img.shape[0]
        width, height = min(w, h * aspect), min(h, w / aspect)
        x += (w - width) / 2; y += (h - height) / 2
        self.slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(width), height=Inches(height))
        self.ax.imshow(img, extent=(x, x + width, y + height, y), aspect="auto")

    def cards(self, cards):
        two_rows = len(cards) > 2
        for i, (heading, body) in enumerate(cards):
            if two_rows:
                col, row = i % 2, i // 2
                x, y, w, h = .65 + col * 6.12, 1.98 + row * 2.12, 5.92, 1.88
                heading_size, body_size = 17, 15
            else:
                x, y, w, h = .65 + i * 6.12, 2.04, 5.92, 3.65
                heading_size, body_size = 20, 18
            self.rect(x, y, w, h)
            self.rect(x, y, .045, h, CYAN)
            self.text(heading, x+.22, y+.19, w-.44, .48, heading_size, CYAN, True)
            self.text(body, x+.22, y+.80 if not two_rows else y+.66,
                      w-.44, h-.96 if not two_rows else h-.76, body_size, MUTED)

    def close(self, takeaway, notes, preview=None):
        self.text(takeaway, .67, 6.53, 11.95, .43, 12, CYAN)
        self.slide.notes_slide.notes_text_frame.text = notes
        self.pdf.savefig(self.fig, facecolor=self.fig.get_facecolor())
        if preview:
            self.fig.savefig(preview, dpi=115, facecolor=self.fig.get_facecolor())
        plt.close(self.fig)


def create_presentation(summary, out):
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    prs.core_properties.title = "Insurance Behavior Lab — EDA to Machine Learning"
    prs.core_properties.subject = "Studi kasus sintetis churn, renewal dan akuisisi asuransi"
    prs.core_properties.author = "Insurance Analytics Workshop"
    figs = out / "reports/figures"
    metrics, acquisition = summary["churn_test"], summary["acquisition_test"]
    slides = []

    def cards(title, section, items, takeaway, notes):
        slides.append((title, section, "cards", items, takeaway, notes))

    def chart(title, section, filename, takeaway, notes):
        slides.append((title, section, "image", filename, takeaway, notes))

    cards("Ketika perilaku nasabah sulit dijelaskan", "Insurance Behavior Lab / Masterclass", [
        ("Misteri bisnis", "Premi naik 10%, nasabah tetap bertahan. Beberapa waktu kemudian churn melonjak. Campaign sepi, lalu pendaftaran ramai."),
        ("Tugas data scientist", "Ubah kebingungan menjadi hipotesis, sinyal terukur, model yang diuji, dan keputusan yang dapat dipertanggungjawabkan.")],
        "Paket pembelajaran dengan data sintetis, kode, notebook, visual dan model yang dapat dijalankan.",
        "Pembukaan 5 menit. Minta peserta membedakan tiga pertanyaan: siapa akan churn, kapan permintaan berubah, dan apa penyebabnya. Semua data dibuat untuk pembelajaran. Tidak ada klaim hubungan dengan perusahaan teknologi tertentu atau hasil bisnis nyata. Sebelum menampilkan model, minta dua hipotesis per kelompok. Tekankan bahwa senioritas analisis terlihat dari definisi target, kualitas bukti dan keputusan yang dapat diuji.")
    cards("Tiga fenomena, tiga pertanyaan berbeda", "01 / Business framing", [
        ("Harga dan ketahanan", "Apakah nasabah menunda keputusan sampai renewal? Apakah nilai proteksi masih tinggi?"),
        ("Churn tertunda", "Apakah tunggakan dan keluhan meningkat sebelum keputusan berhenti?"),
        ("Akuisisi tak terduga", "Apakah respons campaign tertunda, minat organik meningkat, atau attribution belum lengkap?"),
        ("Jenis jawaban", "Deskripsi menjelaskan apa; prediksi memperkirakan risiko; eksperimen menguji dampak.")],
        "Jangan memaksa satu model menjawab semua pertanyaan.",
        "Latihan 10 menit. Buat tabel fenomena, hipotesis, data yang dibutuhkan dan cara menguji. Tetap bertahan bisa timbul akibat friksi pindah, kebutuhan proteksi, autodebit atau jadwal kontrak. Tidak ada respons langsung tidak sama dengan efek harga nol; terlebih tanpa kontrol dan ukuran ketidakpastian yang tepat.")
    chart("Learning path yang menghasilkan artefak", "02 / Roadmap 16 jam", "learning_path",
          "Business brief → dataset → EDA → model → evaluasi → rekomendasi dan eksperimen.",
          "Empat sesi masing-masing empat jam. Sesi pertama framing dan EDA; sesi kedua fitur dan baseline; sesi ketiga model dan evaluasi; sesi keempat interpretasi, akuisisi, skenario dan presentasi. Peserta harus menyerahkan artefak pada setiap tahap. Detail waktu dan rubrik terdapat di docs/LEARNING_PATH.md.")
    cards("Persiapan peserta dan hasil belajar", "02 / Kompetensi", [
        ("Prasyarat", "Python dasar, fungsi, DataFrame, groupby, persentase, rata-rata dan grafik."),
        ("Kemampuan analisis", "Audit snapshot, denominator churn, cohort, interaksi, lag dan censoring."),
        ("Kemampuan ML", "Pipeline preprocessing, baseline, validation waktu, AP, lift dan calibration."),
        ("Komunikasi", "Pisahkan fakta, hipotesis, asumsi bisnis dan bukti kausal dalam memo keputusan.")],
        "GPU tidak diperlukan; kedisiplinan data dan evaluasi lebih penting daripada kompleksitas model.",
        "Lakukan pretest singkat: hitung churn dari 30 pelanggan keluar dengan 1000 aktif; bedakan jumlah dan rate; tunjukkan mengapa accuracy 97% bisa diperoleh tanpa menangkap satu pun churn. Remediasi diberikan lewat notebook pandas sebelum melatih model.")
    cards("Library yang dipasang dan kegunaannya", "03 / Environment", [
        ("Data dan visual", "numpy: simulasi; pandas: tabel dan waktu; matplotlib: grafik, diagram dan PDF."),
        ("Machine learning", "scikit-learn: pipeline, imputasi, encoder, model dan metrik; scipy: komputasi numerik."),
        ("Artifact dan materi", "joblib: simpan model; python-pptx: deck; nbformat, nbclient, ipykernel: notebook."),
        ("Reproducibility", "requirements.txt dan requirements-lock.txt; seed, manifest, pengujian dan versi runtime.")],
        "Mulai dengan python build_case.py; semua output mengikuti direktori proyek.",
        "Tunjukkan setup virtual environment di README. Lingkungan ini diuji dengan Python dan versi library yang tercatat pada manifest. Peserta memakai requirements-lock.txt bila versi Python kompatibel; gunakan requirements.txt untuk instalasi lintas versi yang didukung. Dependensi advanced seperti SHAP, survival dan experiment tracking dijelaskan sebagai pengembangan, tidak dibutuhkan untuk tutorial inti.")
    cards("Kontrak bisnis dan label yang jelas", "04 / Target definition", [
        ("Unit churn", "Satu nasabah = satu polis. Satu baris = polis aktif pada awal suatu bulan."),
        ("Target", "churn_next_month: terminal lapse/cancel selama bulan sesudah snapshot, sebelum label_end_date."),
        ("Unit akuisisi", "Satu baris = wilayah x minggu. Target = jumlah polis baru dalam minggu prediksi."),
        ("Simplifikasi", "Tidak ada reinstatement, multi-policy, grace period atau competing risks dalam simulator.")],
        "Horizon adalah satu bulan kalender, bukan tepat 30 hari.",
        "Contoh: snapshot 1 April memakai fitur sampai 31 Maret. Target adalah kejadian selama 1 April sampai sebelum 1 Mei. Label baru lengkap pada 1 Mei. Karena bulan berbeda panjang, jangan menyebutnya 30 hari secara persis. Definisi lapse nyata perlu disepakati dengan tim aktuaria dan operasional sebelum proyek produksi.")
    chart("Peta hipotesis dan faktor tersembunyi", "04 / Scientific reasoning", "hypothesis_map",
          "Panah menggambarkan dugaan jalur; data observasional belum membuktikan jalur tersebut.",
          "Diskusikan confounding: tekanan ekonomi dapat memengaruhi kemampuan bayar dan pencarian proteksi. Premi naik serentak dengan kalender sehingga fitur premium_change_pct tidak dapat memisahkan efek harga dari semua kejadian waktu. Tanyakan data apa yang bisa mematahkan hipotesis, bukan hanya mendukungnya. Lihat pedoman pengajar setelah peserta mengerjakan EDA.")
    cards("Tabel dan timestamp yang harus diaudit", "05 / Data quality", [
        ("Kunci unik", "customer_id + snapshot_date. Duplikasi dapat menggandakan denominator dan memberi bobot palsu."),
        ("Ketersediaan fitur", "feature_as_of < snapshot_date < label_end_date. Data setelah churn tidak boleh muncul."),
        ("Kelengkapan", "Null pada app_sessions dan region disengaja untuk mengajarkan imputasi yang fit di train."),
        ("Indeks ekonomi", "economic_stress_index sepenuhnya sintetis; data nyata harus memakai tanggal publikasi asli.")],
        "Audit gagal harus menghentikan pipeline sebelum training.",
        "Jalankan audit_policy_data dan baca artifacts/data_audit.json. Pemeriksaan timestamp adalah kontrak data, bukan jaminan bahwa setiap sumber benar. Dalam produksi diperlukan audit event_time, ingestion_time, dan tanggal rilis setiap sumber. Revisi data makro dapat menimbulkan look-ahead bias bila nilai terbaru dipakai untuk snapshot lama.")
    chart("EDA awal: missingness dan distribusi", "05 / Audit visual", "data_quality",
          "Nilai kosong adalah kondisi yang harus dipahami sebelum diisi.",
          "Tanyakan apakah missingness acak atau terkait channel/segmen. Dalam simulasi missingness dibuat acak. Median dan most_frequent dipelajari dari train di dalam pipeline. Jangan fit imputer pada seluruh dataset lalu melakukan split. Distribusi tenure memuat pelanggan lama pada awal observasi dan cohort baru yang masuk kemudian.")
    chart("Premi naik; kapan churn benar-benar bergerak?", "06 / EDA timeline", "churn_trend",
          "Bandingkan rate dan count sekaligus; basis nasabah berubah sepanjang waktu.",
          "Baca tiga periode sebelum harga, segera setelah harga dan periode lonjakan. Grafik bukan uji kausal dan tidak menyatakan tidak signifikan secara statistik. Jelaskan survivorship: orang yang sudah churn tidak bisa muncul lagi, sehingga komposisi survivor berubah. Bedakan churn musiman, perubahan administrasi label, dan perubahan perilaku nyata.")
    chart("Pola interaksi lebih berguna dari rata-rata", "06 / Segmentation", "renewal_payment",
          "Tunggakan dekat renewal dapat menjadi sinyal yang lebih kuat daripada premi saja.",
          "Periksa jumlah sampel setiap segmen. Lanjutkan dengan membandingkan pola di bulan yang berbeda dan income_band. Segmentasi ini deskriptif; nasabah dengan tunggakan tidak identik dengan nasabah tanpa tunggakan. Hindari menyimpulkan bahwa menghilangkan tunggakan secara otomatis akan menghasilkan penurunan churn sesuai selisih rate.")
    chart("Cohort retention dan right censoring", "06 / Cohort analysis", "cohort_retention",
          "Cohort baru dipisahkan dari pelanggan lama yang telah bertahan sebelum data dimulai.",
          "Heatmap memakai product-limit dari rate churn per tenure. Hanya inception cohort digunakan. Sel yang semua anggota cohort belum punya kesempatan follow-up lengkap disembunyikan. Sel kosong bukan retention nol. Kolom pertama adalah bertahan di akhir bulan pertama, bukan snapshot inception. Jelaskan left truncation pada cohort lama dan mengapa mereka tidak boleh disamakan dengan cohort baru.")
    chart("Sinyal sebelum outcome: keluhan dan pembayaran", "06 / Leading indicators", "leading_signals",
          "Semua sinyal tersedia pada snapshot; nilai z-score hanya membantu perbandingan bentuk seri.",
          "Jangan bandingkan besar z-score seolah satuannya sama secara bisnis. Agregat bulanan dapat dipengaruhi komposisi pelanggan. Untuk produksi, bentuk window 30/90 hari dari event timestamp dan simpan lineage fitur. Untuk challenge, bandingkan effect lag dengan model dan segmen, tanpa memilih lag berdasarkan final test.")
    chart("Campaign tidak selalu ramai pada minggu yang sama", "07 / Acquisition EDA", "acquisition_trend",
          "Kebutuhan proteksi, pencarian organik, referral dan jeda respons adalah hipotesis terukur.",
          "Acquisition dataset adalah simulasi analitis terpisah dari cohort polis; bukan buku besar yang direkonsiliasi satu-ke-satu. Tunjukkan kenaikan lead/referral dan kapan polis baru ikut naik. Kolom inbound_leads adalah activity minggu tersebut, bukan conversion denominator cohort yang matang. Jangan menyebut rasio new_policies/inbound_leads sebagai conversion cohort sebenarnya.")
    chart("Menguji dugaan respons campaign yang tertunda", "07 / Lag analysis", "campaign_lag",
          "Korelasi pada lag membantu memilih investigasi; periodicity dan confounder tetap harus diperiksa.",
          "Lag 2 berarti campaign dua minggu sebelumnya dibandingkan penjualan sekarang. Campaign dijadwalkan periodik, sehingga korelasi beberapa lag dapat terbentuk tanpa bukti kausal. Jika menguji banyak lag pada data nyata, gunakan holdout, kontrol kalender dan desain eksperimen; jangan memilih p-value terendah sebagai jawaban.")
    cards("Feature engineering dengan kontrak waktu", "08 / Feature design", [
        ("Riwayat 3 bulan", "late_payments_3m, complaints_3m, claims_3m dihitung sebelum snapshot."),
        ("Konteks renewal", "months_to_renewal dan interaksi renewal x late_payments_3m."),
        ("Keterjangkauan", "premium_to_income = premi bulanan / pendapatan simulasi bulanan."),
        ("Larangan leakage", "ID, churn outcome, tanggal akhir label dan variabel shock tersembunyi bukan fitur.")],
        "Dalam implementasi nyata, historical windows harus dibangun dari event log yang tersedia saat itu.",
        "Simulator menghasilkan agregat historis secara langsung sebelum outcome draw. Dua fitur turunan berada di engineer_policy_features. Ini tidak menggantikan pembangunan ETL dari transaksi pembayaran nyata. Aktivitas sebelum inception dikeluarkan dari rolling counts. Pendapatan dalam produksi mungkin tidak tersedia atau tidak andal; uji model tanpa fitur ini dan evaluasi dampak per segmen.")
    cards("Train, validation, test dan embargo", "09 / Temporal validation", [
        ("Training", "Jan 2023–Jun 2025. Semua preprocessing hanya fit di periode ini."),
        ("Validation", "Agu–Des 2025. Pilih model menggunakan average precision; tentukan threshold."),
        ("Final test", "Feb–Sep 2026. Jangan ubah model setelah melihat hasilnya."),
        ("Embargo", "Jul 2025 dan Jan 2026. Label train/validation matang sebelum scoring split berikutnya.")],
        "Nasabah berulang mencerminkan scoring portofolio; generalisasi ke nasabah baru perlu evaluasi terpisah.",
        "Random split membocorkan informasi kalender dan kondisi cohort. Pada panel, TimeSeriesSplit per baris juga dapat memisahkan wilayah atau pelanggan pada waktu yang sama. Paket ini membagi berdasarkan kalender eksplisit. Data embargo tidak dipakai untuk fit model, tetapi dapat menjadi riwayat fitur yang memang sudah tersedia pada snapshot selanjutnya.")
    cards("Rolling backtest untuk memeriksa stabilitas", "09 / Three historical folds", [
        ("Fold 1", "Train sampai Jun 2024; validasi Agu–Okt 2024."),
        ("Fold 2", "Train sampai Sep 2024; validasi Nov 2024–Jan 2025."),
        ("Fold 3", "Train sampai Des 2024; validasi Feb–Apr 2025."),
        ("Interpretasi", "Bandingkan AP, lift dan prevalensi antar fold; hindari kesimpulan dari satu periode mudah.")],
        "Lihat artifacts/tables/rolling_backtest.csv; ketiga model memakai kontrak fitur yang sama.",
        "Parameter kandidat ditentukan sebelum validation dan test. Rolling backtest dipakai sebagai diagnostik historis; pilihan final tetap berdasarkan validation utama. Fold yang belum pernah melihat kenaikan harga tidak dapat belajar efek harga dari data training. Ini contoh nyata distribution shift dan keterbatasan observasi, bukan alasan untuk menyisipkan outcome masa depan.")
    chart("Pipeline ML yang dapat dipakai ulang", "10 / Engineering workflow", "ml_process",
          "Preprocessing dan estimator disimpan bersama agar training dan inference konsisten.",
          "Buka modeling.py. ColumnTransformer memisahkan numerik dan kategori. SimpleImputer menangani missing values, StandardScaler membantu logistik dan Poisson, OneHotEncoder menangani kategori tak dikenal. Pipeline disimpan dengan joblib. Jangan memuat artifact joblib dari sumber tak dipercaya. Infrastruktur CRM, feature store dan monitoring dalam diagram adalah rancangan berikutnya, bukan integrasi produksi yang sudah aktif.")
    cards("Baseline dahulu, lalu model nonlinear", "10 / Churn models", [
        ("Baseline prevalensi", "Semua nasabah mendapat probabilitas churn training yang sama. Menunjukkan nilai tambah ranking."),
        ("Logistic regression", "Probabilitas dari kombinasi fitur linear pada log-odds; baseline yang mudah diaudit."),
        ("Gradient boosting", "Pohon kecil ditambahkan bertahap untuk menangkap hubungan nonlinear dan interaksi."),
        ("Pilihan sadar", "Tidak memakai class weighting otomatis karena probabilitas dan kalibrasi penting untuk keputusan.")],
        "Model rumit harus mengalahkan baseline pada data masa depan, bukan hanya training.",
        "Jelaskan logit p = beta0 + beta x dan sigmoid p = 1/(1+exp(-z)). HGB menggunakan max_leaf_nodes=15, min_samples_leaf=55, L2=8 dan 130 iterasi. Early stopping random internal dimatikan agar tidak mengganti strategi waktu. Hyperparameter fixed bukan hasil pencarian final test. GPU dan deep learning tidak diperlukan untuk ukuran dan jenis data ini.")
    cards("Metrik yang relevan dengan churn langka", "11 / Evaluation", [
        ("Average precision", "Merangkum precision-recall berbobot kenaikan recall. Baseline ranking acak sekitar prevalensi."),
        ("ROC-AUC dan Brier", "ROC-AUC mengukur ranking pasangan; Brier mengukur squared error probabilitas."),
        ("Capacity metrics", "Precision, recall dan lift pada top 10% risiko yang dihitung terpisah setiap bulan."),
        ("Kesalahan umum", "Accuracy tinggi dapat berasal dari memprediksi semua nasabah bertahan.")],
        "AP di paket ini bukan trapezoidal PR-AUC; gunakan istilah metrik dengan tepat.",
        "Contoh 1000 nasabah dan 30 churn: prediksi semua bertahan memberi accuracy 97% tetapi recall churn nol. Precision adalah true positive dibagi semua yang dihubungi; recall adalah true positive dibagi seluruh churn; lift membandingkan precision dengan prevalensi. Skor probabilitas perlu diperiksa dengan reliability diagram dan Brier sebelum dipakai untuk nilai harapan.")
    chart("Perbandingan model pada final test", "11 / Results from this run", "model_comparison",
          f"Terpilih dari validation: {summary['churn_model']} | AP test {metrics['average_precision']:.3f}.",
          "Angka di slide ini diambil otomatis dari run_summary.json. Perhatikan model terpilih ditentukan oleh validation, sehingga hasil test tidak selalu menjadi model dengan nilai tertinggi di antara semua kandidat. Itu perilaku evaluasi yang sehat. Jika hasil test buruk, laporkan sebagai temuan dan rancang iterasi dengan holdout baru, bukan memilih model ulang berdasarkan test.")
    chart("Precision-recall dan ROC membaca hal berbeda", "11 / Ranking diagnostics", "pr_roc",
          f"AP test {metrics['average_precision']:.3f} | ROC-AUC {metrics['roc_auc']:.3f} | prevalensi {metrics['prevalence']:.1%}.",
          "Pada PR curve, precision baseline mendekati prevalensi. Metrik berubah ketika prevalensi berubah, sehingga bandingkan antar waktu dengan konteks base rate. ROC-AUC relatif kurang sensitif pada prevalensi namun tidak menjawab kapasitas call center. Jangan mengubah threshold setelah melihat final test untuk mempercantik confusion matrix.")
    chart("Apakah probabilitas dapat dipercaya?", "11 / Calibration diagnostic", "calibration",
          f"Brier test {metrics['brier']:.4f}; grafik adalah diagnostik, bukan klaim telah melakukan recalibration.",
          "Kelompok dengan rata-rata prediksi 0.2 seharusnya memiliki churn sekitar 20% bila model terkalibrasi. Deviasi dari diagonal dapat berasal dari distribution shift, model misspecification atau data sedikit. Bila ingin melakukan sigmoid/isotonic calibration, siapkan periode kalibrasi khusus atau cross-validation temporal di development data, dan evaluasi lagi pada holdout baru.")
    chart("Ranking bulanan sesuai kapasitas kontak", "12 / Operational evaluation", "decile_lift",
          f"Top 10%: precision {metrics['precision_at_10pct']:.1%}, recall {metrics['recall_at_10pct']:.1%}, lift {metrics['lift_at_10pct']:.2f}x.",
          "Ranking dilakukan setiap bulan untuk menghindari mengalokasikan kapasitas periode sekarang memakai skor dari masa depan. ceil memastikan sedikitnya satu orang ketika dataset kecil. Ikatan skor dipecahkan dengan customer_id secara deterministik; pada baseline prevalensi ini hanya aturan acak semu, bukan kemampuan prediksi.")
    chart("Threshold tetap dan top-K berbeda", "12 / Error analysis", "confusion_matrix",
          f"Threshold validation = {summary['validation_threshold']:.4f}; jumlah kandidat per bulan bisa berubah.",
          "Threshold tetap memakai quantile 90% skor validation. Ini bukan aturan top10 bulanan yang dipakai metrik capacity. Analisis false positive: apakah berisiko tetapi tidak churn karena faktor lain? False negative: faktor apa belum diobservasi? Buka churn_test_predictions.csv dan pilih masing-masing dua kasus untuk diskusi tanpa menganggap model mengetahui alasan sebenarnya.")
    cards("Ketidakpastian dan nasabah yang berulang", "12 / Statistical limits", [
        ("Bootstrap nasabah", "Sampling ulang ID nasabah mempertahankan seluruh snapshot milik ID tersebut."),
        ("Interval AP", f"95%: {summary['customer_bootstrap_intervals']['average_precision']['lower_95']:.3f}–{summary['customer_bootstrap_intervals']['average_precision']['upper_95']:.3f}."),
        ("Interval lift", f"95%: {summary['customer_bootstrap_intervals']['lift_at_10pct']['lower_95']:.2f}–{summary['customer_bootstrap_intervals']['lift_at_10pct']['upper_95']:.2f}x."),
        ("Batas interpretasi", "Model tetap; shock waktu bersama dan seluruh ketidakpastian seleksi model belum tercakup.")],
        "120 pengulangan adalah ilustrasi praktikum, bukan jaminan coverage pada semua proses bisnis.",
        "Bootstrap baris menganggap snapshot orang yang sama independen. Bootstrap customer mengurangi masalah tersebut, tetapi customer dalam bulan sama masih terkena shock bersama. Analisis lanjutan dapat memakai block bootstrap waktu atau evaluasi beberapa periode independen. Interval tidak boleh disalahartikan sebagai probabilitas bahwa efek kampanye bersifat kausal.")
    chart("Apa yang digunakan model untuk meranking?", "13 / Explainability", "feature_importance",
          "Importance dihitung pada validation, sehingga final test tidak dipakai untuk memilih fitur.",
          "Permutation importance mengacak satu kolom lalu mengukur penurunan AP. Error bar adalah standar deviasi empat pengulangan, bukan confidence interval kausal. Fitur berkorelasi dapat menyebabkan importance dibagi atau tampak kecil. Mengacak fitur yang saling bergantung juga dapat membuat kombinasi tak realistis. Gunakan hasilnya sebagai diagnosis model, bukan instruksi mengubah kebijakan harga.")
    cards("Menjawab misteri tanpa klaim berlebihan", "13 / Evidence and hypotheses", [
        ("Bertahan setelah harga naik", "Periksa inertia, autodebit, nilai proteksi, renewal timing dan perubahan mix nasabah."),
        ("Churn meningkat kemudian", "Cari kenaikan tunggakan, keluhan, cluster renewal dan perubahan pencatatan lapse."),
        ("Akuisisi meningkat", "Cari minat pencarian, referral, berita proteksi, channel dan jeda campaign."),
        ("Data tambahan", "Exposure per nasabah, alasan cancel, grace period, kompetitor dan timestamp rilis makro.")],
        "Tidak signifikan secara statistik membutuhkan desain uji; grafik yang tampak datar belum cukup.",
        "Hasil pembelajaran terbaik adalah daftar hipotesis yang bisa dibantah. Pilih satu hipotesis, tentukan outcome dan kontrol, lalu jelaskan pengamatan apa yang akan membuat tim meninggalkan hipotesis itu. Tindakan meningkatkan layanan mungkin tetap layak diuji walaupun mekanisme churn belum diketahui secara lengkap.")
    cards("Akuisisi memerlukan model sendiri", "14 / Weekly demand model", [
        ("Baseline lag-1", "Perkiraan minggu ini = polis baru minggu lalu, per wilayah."),
        ("Poisson regression", "Model mean hitungan nonnegatif menggunakan log link dan regularisasi."),
        ("Fitur sebelum minggu", "Lag polis, lead, pencarian, referral; campaign terencana dan campaign dua minggu lalu."),
        ("Evaluasi", "Split berdasarkan minggu untuk seluruh wilayah sekaligus; MAE total dan WAPE.")],
        "Actual minggu sebelumnya tersedia di setiap origin; ini backtest satu minggu ke depan.",
        "Jelaskan distribusi Poisson sebagai model hitungan dengan log mean. Dalam data nyata overdispersion dan zero inflation mungkin perlu model lain. MAE region-week berbeda dari MAE penjumlahan mingguan, karena error antarwilayah bisa saling mengimbangi. Paket memilih model berdasarkan MAE total validation sesuai keputusan kapasitas nasional.")
    chart("Prediksi akuisisi dibandingkan actual", "14 / Acquisition results", "acquisition_forecast",
          f"{summary['acquisition_model']} | MAE total {acquisition['weekly_total_mae']:.1f} polis/minggu | WAPE {acquisition['weekly_total_wape']:.1%}.",
          "Skor yang baik tidak mengungkap sebab minat organik. Tanyakan apakah baseline lag1 cukup kuat dan kapan model lebih lambat merespons perubahan. Jangan menyebut garis ini forecast delapan bulan yang dibuat sekali pada Februari: setiap titik menggunakan observed lag dari minggu sebelumnya. Untuk forecast multi-step, perlu recursive/direct forecasting tanpa actual masa depan.")
    cards("Residual menjadi antrean investigasi", "14 / Anomaly triage", [
        ("Definisi", "Residual = actual polis baru - prediksi pada wilayah dan minggu yang sama."),
        ("Aturan praktikum", "Flag di luar quantile 2.5%–97.5% residual validation."),
        ("Investigasi", "Cek tracking, event lokal, kanal baru, referral, pesaing dan perubahan underwriting."),
        ("Batas", "Aturan residual bukan interval prediksi terkalibrasi; rate false alarm perlu dipantau.")],
        "Anomali adalah pertanyaan operasional yang diprioritaskan, bukan penjelasan otomatis.",
        "Lihat acquisition_test_predictions.csv. Flag positif dapat berarti permintaan tak terduga, pencatatan backlog atau kesalahan ETL. Flag negatif bisa berarti gangguan channel atau kualitas lead. Pelajari residual per wilayah karena skala volume berbeda. Jika ingin jaminan coverage, pelajari conformal forecasting dengan asumsi temporal yang eksplisit sebagai materi advanced.")
    cards("Dari skor risiko ke tindakan layanan", "15 / Retention playbook", [
        ("Keluhan", "Tinjau keluhan belum selesai dan uji service recovery yang relevan."),
        ("Pembayaran", "Tawarkan bantuan dan penjelasan opsi pembayaran sesuai kebijakan perusahaan."),
        ("Renewal", "Berikan pengingat dan penjelasan manfaat sebelum keputusan pembaruan."),
        ("Prioritas", "Gabungkan risiko, nilai, kapasitas dan respons intervensi yang terukur.")],
        "Risiko churn tinggi tidak otomatis berarti nasabah mudah diselamatkan.",
        "Suggested_review pada tabel prediksi adalah aturan sederhana yang transparan, bukan treatment recommendation teroptimasi. Intervensi harus dipilih dengan tim operasional dan dievaluasi menggunakan randomisasi. Jangan menyamakan tinggi risiko dengan pelanggan layak ditolak; tujuan tutorial adalah layanan dan retensi yang bermanfaat.")
    chart("Nilai bisnis bergantung pada efek yang belum diukur", "15 / Decision economics", "business_sensitivity",
          "E[net] = P(churn) x rescue_fraction x margin_terselamatkan - biaya kontak.",
          "Margin asumsi adalah 20% dari premi 12 bulan. Biaya kontak Rp15000 dibebankan ke semua nasabah yang dihubungi. Rescue_fraction adalah peluang menyelamatkan seseorang yang akan churn tanpa intervensi, bukan treatment effect terukur. Formula menyederhanakan potensi dampak negatif, discount cost, dan horizon. Kurva sensitivitas membantu menilai kebutuhan bukti sebelum rollout.")
    cards("Eksperimen retensi yang dapat menjawab dampak", "16 / Causal experiment", [
        ("Eligibility", "Tetapkan kelompok eligible dan aturan kontak sebelum randomisasi."),
        ("Randomisasi", "Treatment dan control pada nasabah, stratifikasi risiko/wilayah bila diperlukan."),
        ("Outcome utama", "Selisih retention setelah horizon; intention-to-treat; biaya dan margin incremental."),
        ("Guardrail", "Keluhan, beban kontak, opt-out, spillover dan fairness per kelompok.")],
        "Ukuran sampel memerlukan baseline, minimum detectable effect, power dan unit randomisasi.",
        "Contoh: bandingkan pengingat plus bantuan layanan dengan layanan rutin pada high-risk eligible. Jangan memilih control dari nasabah low-risk. Definisikan kapan outcome matang, hindari menghentikan eksperimen saat hasil kebetulan baik, dan tentukan analisis sebelum outcome dibuka. Power calculation harus mempertimbangkan churn langka dan cluster bila randomisasi berbasis wilayah.")
    cards("Eksperimen campaign yang memisahkan incrementality", "16 / Acquisition experiment", [
        ("Pertanyaan", "Berapa polis tambahan yang tercipta dibanding tidak menjalankan campaign?"),
        ("Desain", "Randomized geo/channel holdout dengan wilayah sebanding dan durasi mencakup lag."),
        ("Ukuran sukses", "Incremental policies dan margin setelah biaya; attribution sendiri tidak cukup."),
        ("Risiko desain", "Spillover lintaswilayah, seasonality, overlap campaign dan perubahan channel.")],
        "Jika randomisasi tidak tersedia, quasi-experiment membutuhkan asumsi yang harus diuji.",
        "Difference-in-differences memerlukan counterfactual yang masuk akal dan asumsi parallel trends, bukan sekadar membandingkan before-after. Jangan mengklaim feature campaign dalam model Poisson sebagai pengukuran ROI kampanye. Tahap pertama adalah exposure logging yang dapat dipercaya, diikuti desain eksperimen dan power.")
    cards("Model artifact dan batch scoring", "17 / Reproducible engineering", [
        ("Training", "python build_case.py --customers 6000 --seed 26"),
        ("Artifact", "artifacts/models/churn_model.joblib berisi pipeline, fitur, threshold dan metadata."),
        ("Scoring", "python -m insurance_lab.score dengan --model, --input, --output."),
        ("Review", "Gunakan scoring_example.csv; lihat hasil batch sebelum integrasi sistem lain.")],
        "README memuat command lengkap, termasuk jalur input dan output.",
        "Model disimpan bersama seluruh preprocessing, sehingga pengguna tidak menulis ulang imputasi atau mapping kategori saat scoring. Schema fitur diverifikasi sebelum predict_proba. Contoh inference adalah snapshot historis sintetis, bukan daftar nasabah live. Artifact joblib hanya boleh berasal dari sumber yang dipercaya karena deserialisasi dapat menjalankan kode.")
    chart("Distribution shift harus ditinjau", "18 / Monitoring diagnostics", "distribution_shift",
          "Perbedaan distribusi tidak otomatis berarti model gagal; hubungkan dengan kinerja dan konteks.",
          "Grafik memakai standardized mean difference, bukan PSI. Fitur yang variance train nol dikeluarkan dari rasio. Drift dapat disebabkan seleksi survivor, perubahan harga, campuran wilayah atau kualitas data. Gunakan metrik bulanan dan subgroup_diagnostics.csv sebagai pemeriksaan pendamping. Ambang alert perlu ditentukan dari toleransi operasional dan data historis.")
    cards("Monitoring dan tata kelola setelah tutorial", "18 / Production design", [
        ("Data", "Schema, null, duplikasi, ketepatan timestamp, freshness dan missing channel."),
        ("Model", "AP, lift, Brier, reliability dan stabilitas per bulan serta kelompok."),
        ("Bisnis", "Incremental retention, biaya, margin, conversion dan kepuasan nasabah."),
        ("Operasi", "Akses data, audit log, review manusia, rollback dan jadwal evaluasi ulang.")],
        "Dashboard tutorial adalah laporan offline; monitoring produksi membutuhkan integrasi dan outcome yang matang.",
        "Perlu pemilik model, pemilik data dan pemilik keputusan. Retrain tidak boleh otomatis hanya karena drift alert. Periksa dulu apakah terjadi masalah ETL, perubahan label atau kondisi bisnis. Data pribadi nyata harus dikelola sesuai tata kelola organisasi dan ketentuan yang berlaku; paket ini hanya memakai data sintetis tanpa identitas nyata.")
    chart("Dashboard untuk menjelaskan keputusan", "19 / Executive communication", "executive_dashboard",
          "Presentasikan pola, kualitas prediksi, asumsi dan eksperimen berikutnya dalam satu cerita.",
          "Gunakan narasi: apa yang terlihat, apa yang dapat diprediksi, apa yang belum diketahui, dan apa yang akan diuji. Jangan menampilkan semua metrik tanpa keputusan. Laporan dashboard.html menyertakan gambar agar bisa dibuka offline; INSIGHTS.md menyimpan narasi detail dan batas analisis. Semua angka dibuat ulang dari hasil pipeline.")
    cards("Praktikum: deliverable per kelompok", "20 / Student lab", [
        ("Lab 1: framing dan EDA", "Tiga hipotesis; audit; lima grafik; memo yang membedakan fakta dan dugaan."),
        ("Lab 2: fitur dan baseline", "Definisi label; daftar leakage; split waktu; reproduksi baseline dan logistik."),
        ("Lab 3: evaluasi", "Bandingkan AP/lift; bahas calibration; tinjau false positive dan false negative."),
        ("Lab 4: keputusan", "Bandingkan forecast; sensitivitas biaya; desain eksperimen; pitch lima menit.")],
        "Notebook dapat dijalankan per tahap; jangan membuka jawaban simulator sebelum hipotesis disusun.",
        "Notebook utama menjalankan tahap inti dan memberi pertanyaan setelah output. Notebook challenge berisi tugas tanpa jawaban implementasi. Pengajar menilai proses dan kualitas bukti, bukan mencari AUC terbesar. Wajib laporkan bila model nonlinear kalah terhadap logistik atau baseline; hasil itu dapat menjadi temuan yang baik jika analisisnya benar.")
    cards("Rubrik dan pertanyaan penilaian", "20 / Assessment", [
        ("Business framing: 20%", "Target, unit, definisi sukses dan hipotesis yang bisa diuji."),
        ("Data dan EDA: 25%", "Timestamp, denominator, cohort, missingness dan visual yang tepat."),
        ("Model dan validasi: 30%", "Baseline, temporal split, metrik, batas generalisasi dan reproducibility."),
        ("Keputusan: 25%", "Insight 15%; eksperimen dan komunikasi 10%. Hindari klaim kausal tanpa desain.")],
        "Reproducible dan jujur tentang bukti merupakan syarat dasar semua komponen.",
        "Pertanyaan lisan: mengapa accuracy tidak cukup; kapan label matang; mengapa random split bermasalah; apakah feature importance membuktikan penyebab; apakah forecast satu langkah sama dengan banyak langkah; bagaimana mengukur manfaat retention. Beri kredit untuk hipotesis yang dibatalkan oleh data dan untuk menjelaskan keterbatasan secara tepat.")
    cards("Pengembangan berikutnya sesuai kebutuhan data", "21 / Advanced track", [
        ("Survival analysis", "Model waktu menuju lapse dengan right censoring dan covariates berbasis waktu."),
        ("Causal dan uplift", "Pelajari heterogenitas respons setelah eksperimen treatment-control tersedia."),
        ("Explainability", "SHAP dan stability checks; tetap perhatikan fitur berkorelasi dan interpretasi kausal."),
        ("MLOps", "Experiment tracking, feature lineage, CI, drift monitoring dan kontrol deployment.")],
        "Advanced track adalah materi lanjutan; library tersebut tidak diperlukan untuk menjalankan paket inti.",
        "Hindari memasang seluruh ekosistem advanced sebelum pertanyaan bisnis jelas. Survival berguna ketika waktu event dan censoring menjadi pusat masalah. Uplift memerlukan informasi treatment dan asumsi identifikasi. Experiment tracking bermanfaat ketika kandidat dan tim bertambah, tetapi manifest dan kode bersih cukup untuk laboratorium pertama.")
    cards("Sumber belajar dan penutupan kelas", "22 / References and next steps", [
        ("Dokumentasi teknis", "scikit-learn: temporal validation, average precision, calibration dan permutation importance."),
        ("Konteks Indonesia", "BPS: metadata IHK; Bank Indonesia: indikator BI-Rate. Belum digabung ke data sintetis ini.")],
        "Prediksi membantu memilih tindakan; eksperimen membantu mengetahui tindakan mana yang bermanfaat.",
        "URL lengkap dan keterangan penggunaan ada di docs/REFERENCES.md. Periksa tanggal publikasi sebelum menggabungkan data makro ke snapshot historis. Akhiri dengan meminta setiap kelompok menyebutkan satu insight yang didukung data, satu hal yang belum diketahui, dan satu eksperimen yang dapat mengurangi ketidakpastian.")

    with PdfPages(out / "slides/Insurance_Behavior_Lab.pdf") as pdf:
        for number, (title, section, kind, payload, takeaway, notes) in enumerate(slides, 1):
            slide = Slide(prs, number, title, section, pdf)
            if kind == "image":
                slide.image(figs / f"{payload}.png")
            else:
                slide.cards(payload)
            preview = out / f"slides/preview_{number:02d}.png" if number in [1, 3, 10, 22, 38] else None
            slide.close(takeaway, notes, preview)
    prs.save(out / "slides/Insurance_Consumer_Behavior_ML_Masterclass.pptx")
    return len(slides)
