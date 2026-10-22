"""Meaningful checks for leakage, terminal events, temporal lags and saved inference."""
import json
import unittest
from pathlib import Path
import joblib
import numpy as np
import pandas as pd
from pptx import Presentation

from insurance_lab.config import ROOT, TARGET, FEATURES, split_months
from insurance_lab.data import (audit_policy_data, engineer_acquisition_features,
                               engineer_policy_features, simulate_policies)
from insurance_lab.modeling import capacity_selection


class DataContracts(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.raw = simulate_policies(1200, 26)
        cls.frame = engineer_policy_features(cls.raw)

    def test_seed_reproducibility(self):
        pd.testing.assert_frame_equal(self.raw, simulate_policies(1200, 26))

    def test_no_snapshots_after_terminal_churn(self):
        audit = audit_policy_data(self.frame)
        self.assertEqual(audit['customers'], 1200)
        self.assertEqual(audit['duplicate_keys'], 0)

    def test_audit_rejects_future_features(self):
        broken = self.frame.copy()
        broken.loc[0, 'feature_as_of'] = broken.loc[0, 'label_end_date']
        with self.assertRaisesRegex(AssertionError, 'Future feature'):
            audit_policy_data(broken)

    def test_audit_rejects_duplicate_keys(self):
        broken = pd.concat([self.frame, self.frame.iloc[[0]]], ignore_index=True)
        with self.assertRaisesRegex(AssertionError, 'Duplicate'):
            audit_policy_data(broken)

    def test_audit_rejects_snapshot_after_churn(self):
        broken = self.frame.copy()
        ids = broken.customer_id.value_counts()
        identifier = ids.loc[ids > 2].index[0]
        index = broken.loc[broken.customer_id == identifier].index[0]
        broken.loc[index, TARGET] = 1
        with self.assertRaisesRegex(AssertionError, 'after terminal churn'):
            audit_policy_data(broken)

    def test_label_maturity_and_calendar_splits(self):
        splits = split_months(self.frame)
        self.assertLess(splits['train'].label_end_date.max(), splits['validation'].snapshot_date.min())
        self.assertLess(splits['validation'].label_end_date.max(), splits['test'].snapshot_date.min())
        used = pd.concat(splits.values())
        self.assertFalse(used.snapshot_date.isin(pd.to_datetime(['2025-07-01', '2026-01-01'])).any())
        self.assertFalse(set(FEATURES) & {'customer_id', TARGET, 'label_end_date'})

    def test_lags_do_not_cross_regions_or_read_current_outcome(self):
        rows = []
        for region, offset in [('A', 10), ('B', 100)]:
            for i, week in enumerate(pd.date_range('2025-01-06', periods=8, freq='W-MON')):
                rows.append(dict(region=region, week=week, new_policies=offset+i,
                                 inbound_leads=200+i, referrals=20+i, search_index=30+i,
                                 campaign_planned=int(i % 2 == 0)))
        frame = pd.DataFrame(rows)
        built = engineer_acquisition_features(frame)
        self.assertTrue((built.new_policies_lag1 == built.new_policies-1).all())
        self.assertTrue((built.new_policies_lag4 == built.new_policies-4).all())
        # Changing current outcome cannot change same-week historical features.
        changed = frame.copy()
        changed.loc[changed.week == changed.week.max(), 'new_policies'] = 99999
        rebuilt = engineer_acquisition_features(changed)
        cols = [c for c in built if c.endswith('lag1') or c.endswith('lag4')]
        pd.testing.assert_frame_equal(built[cols], rebuilt[cols])

    def test_capacity_is_allocated_per_month(self):
        frame = pd.DataFrame({'snapshot_date': ['2025-01']*20 + ['2025-02']*10,
                              'customer_id': range(30), 'score': np.arange(30)/30})
        selected = capacity_selection(frame, .10)
        counts = selected.groupby(frame.snapshot_date).sum().to_dict()
        self.assertEqual(counts, {'2025-01': 2, '2025-02': 1})


class GeneratedArtifacts(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if not (ROOT / 'artifacts/manifest.json').exists():
            raise unittest.SkipTest('Run python build_case.py first for artifact integration checks')

    def test_saved_pipeline_reproduces_test_predictions(self):
        bundle = joblib.load(ROOT / 'artifacts/models/churn_model.joblib')
        features = pd.read_csv(ROOT / 'data/processed/policy_features.csv')
        predictions = pd.read_csv(ROOT / 'artifacts/tables/churn_test_predictions.csv')
        final = features.loc[features.snapshot_date >= '2026-02-01']
        actual = bundle['pipeline'].predict_proba(final[bundle['features']])[:, 1]
        np.testing.assert_allclose(actual, predictions.score, rtol=1e-10, atol=1e-12)

    def test_selection_uses_validation(self):
        for name, metric, ascending in [('churn', 'average_precision', False),
                                         ('acquisition', 'weekly_total_mae', True)]:
            scores = pd.read_csv(ROOT / f'artifacts/tables/{name}_metrics.csv')
            valid = scores.loc[(scores.split == 'validation') & (scores.model != 'prevalence_baseline')]
            best = valid.sort_values(metric, ascending=ascending).iloc[0].model
            selected = set(scores.loc[scores.selected_on_validation, 'model'])
            self.assertEqual(selected, {best})

    def test_powerpoint_notes_images_and_bounds(self):
        deck = Presentation(ROOT / 'slides/Insurance_Consumer_Behavior_ML_Masterclass.pptx')
        manifest = json.loads((ROOT / 'artifacts/manifest.json').read_text())
        self.assertEqual(len(deck.slides), manifest['slides'])
        self.assertGreaterEqual(len(deck.slides), 40)
        for slide in deck.slides:
            self.assertGreater(len(slide.notes_slide.notes_text_frame.text), 100)
            for shape in slide.shapes:
                self.assertGreaterEqual(shape.left, 0)
                self.assertGreaterEqual(shape.top, 0)
                self.assertLessEqual(shape.left + shape.width, deck.slide_width + 100)
                self.assertLessEqual(shape.top + shape.height, deck.slide_height + 100)

    def test_metric_ranges(self):
        scores = pd.read_csv(ROOT / 'artifacts/tables/churn_metrics.csv')
        for col in ['roc_auc', 'average_precision', 'brier', 'precision_at_10pct', 'recall_at_10pct']:
            self.assertTrue(scores[col].between(0, 1).all(), col)
        self.assertTrue((scores.lift_at_10pct > 0).all())


if __name__ == '__main__':
    unittest.main()
